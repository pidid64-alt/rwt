"""Рендеринг документа отчёта: PDF (reportlab), DOCX (python-docx), XLSX (openpyxl), CSV, JSON."""
from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path

from .model import Doc

FONT_DIRS = ["/usr/share/fonts/truetype/dejavu", "/usr/share/fonts/dejavu", "/usr/share/fonts/TTF", "C:/Windows/Fonts"]
STATUS_COLORS = {"ok": "C8E6C9", "check": "FFF3C4", "fail": "F8C9C9", "missing": "FFD9B0", "out_of_range": "F8C9C9", "na": "E5E7EB"}


def _font_paths():
    for d in FONT_DIRS:
        for regular, bold in (("DejaVuSans.ttf", "DejaVuSans-Bold.ttf"), ("arial.ttf", "arialbd.ttf")):
            r, b = Path(d) / regular, Path(d) / bold
            if r.exists() and b.exists():
                return str(r), str(b)
    return None, None


# ───────────────────────────── PDF ─────────────────────────────
def to_pdf(doc: Doc) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, KeepTogether

    reg, bold = _font_paths()
    if reg:
        pdfmetrics.registerFont(TTFont("F", reg))
        pdfmetrics.registerFont(TTFont("FB", bold))
        f, fb = "F", "FB"
    else:  # pragma: no cover
        f, fb = "Helvetica", "Helvetica-Bold"
    st_n = ParagraphStyle("n", fontName=f, fontSize=8.5, leading=11)
    st_c = ParagraphStyle("c", fontName=f, fontSize=7, leading=8.6)
    st_h = ParagraphStyle("h", fontName=fb, fontSize=7, leading=8.6, textColor=colors.white)
    st_h1 = ParagraphStyle("h1", fontName=fb, fontSize=14, leading=17, spaceBefore=8, spaceAfter=4, textColor=colors.HexColor("#0f2a43"))
    st_h2 = ParagraphStyle("h2", fontName=fb, fontSize=11, leading=14, spaceBefore=7, spaceAfter=3, textColor=colors.HexColor("#0b5cad"))
    st_h3 = ParagraphStyle("h3", fontName=fb, fontSize=9, leading=11, spaceBefore=5, spaceAfter=2)
    st_note = ParagraphStyle("note", fontName=f, fontSize=8, leading=10, textColor=colors.HexColor("#8a4b00"), backColor=colors.HexColor("#FFF7E0"), borderPadding=3, spaceBefore=2, spaceAfter=2)
    st_t = ParagraphStyle("t", fontName=fb, fontSize=18, leading=22, textColor=colors.HexColor("#0f2a43"))
    st_dr = ParagraphStyle("dr", fontName=fb, fontSize=10, leading=13, textColor=colors.HexColor("#b91c1c"), backColor=colors.HexColor("#FDECEC"), borderPadding=5, spaceAfter=4)

    def esc(s):
        return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    buf = io.BytesIO()
    pg = landscape(A4)
    W = pg[0] - 20 * mm

    def footer(canvas, d_):
        canvas.saveState()
        canvas.setFont(f, 7)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(10 * mm, 6 * mm, f"РЗА-АТ • {doc.title} • {doc.subtitle}")
        canvas.drawRightString(pg[0] - 10 * mm, 6 * mm, f"стр. {d_.page}")
        if doc.draft:
            canvas.setFont(fb, 40)
            canvas.setFillColor(colors.Color(0.85, 0.1, 0.1, alpha=0.10))
            canvas.translate(pg[0] / 2, pg[1] / 2)
            canvas.rotate(30)
            canvas.drawCentredString(0, 0, "ПРЕДВАРИТЕЛЬНО — НЕ ДЛЯ ВЫДАЧИ")
        canvas.restoreState()

    d = SimpleDocTemplate(buf, pagesize=pg, leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm, bottomMargin=12 * mm, title=doc.title)
    story = [Paragraph(esc(doc.title), st_t), Paragraph(esc(doc.subtitle), st_n), Spacer(1, 4)]
    if doc.draft:
        story.append(Paragraph("ПРЕДВАРИТЕЛЬНЫЙ ДОКУМЕНТ. Окончательная карта уставок не может быть выдана — полный контроль проекта не пройден:<br/>" + "<br/>".join("• " + esc(r) for r in doc.draft_reasons[:12]), st_dr))
    if doc.meta:
        t = Table([[Paragraph(f"<b>{esc(k)}</b>", st_c), Paragraph(esc(v), st_c)] for k, v in doc.meta], colWidths=[40 * mm, W - 40 * mm])
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#bbbbbb")), ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#eef3f8")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story += [t, Spacer(1, 6)]
    for b in doc.blocks:
        if b.kind == "h1":
            story.append(Paragraph(esc(b.text), st_h1))
        elif b.kind == "h2":
            story.append(Paragraph(esc(b.text), st_h2))
        elif b.kind == "h3":
            story.append(Paragraph(esc(b.text), st_h3))
        elif b.kind == "p":
            story.append(Paragraph(esc(b.text), st_n))
        elif b.kind == "note":
            story.append(Paragraph(esc(b.text), st_note))
        elif b.kind == "pagebreak":
            story.append(PageBreak())
        elif b.kind == "image":
            from reportlab.lib.utils import ImageReader
            ir = ImageReader(io.BytesIO(b.data))
            iw, ih = ir.getSize()
            w = min(W * 0.8, 200 * mm)
            story += [Image(io.BytesIO(b.data), width=w, height=w * ih / iw), Paragraph(esc(b.caption), st_c), Spacer(1, 4)]
        elif b.kind == "table" and b.header:
            n = len(b.header)
            ws = b.widths if len(b.widths) == n else [1.0] * n
            tot = sum(ws)
            colw = [W * w / tot for w in ws]
            data = [[Paragraph(esc(h), st_h) for h in b.header]] + [[Paragraph(esc(c), st_c) for c in r] for r in b.rows]
            t = Table(data, colWidths=colw, repeatRows=1)
            style = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f2a43")), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#c8ced6")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                     ("LEFTPADDING", (0, 0), (-1, -1), 2.5), ("RIGHTPADDING", (0, 0), (-1, -1), 2.5), ("TOPPADDING", (0, 0), (-1, -1), 1.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5)]
            for i, s in enumerate(b.row_status):
                if s in STATUS_COLORS:
                    style.append(("BACKGROUND", (0, i + 1), (-1, i + 1), colors.HexColor("#" + STATUS_COLORS[s])))
                elif i % 2:
                    style.append(("BACKGROUND", (0, i + 1), (-1, i + 1), colors.HexColor("#f6f8fa")))
            t.setStyle(TableStyle(style))
            if b.caption:
                story.append(Paragraph(esc(b.caption), st_c))
            story += [t, Spacer(1, 5)]
    d.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


# ───────────────────────────── DOCX ─────────────────────────────
def _shade(cell, hex_color: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def to_docx(doc: Doc) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Cm, Pt, RGBColor
    d = Document()
    sec = d.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Cm(1.5))
    st = d.styles["Normal"]
    st.font.name = "Calibri"; st.font.size = Pt(9)
    d.add_heading(doc.title, 0)
    d.add_paragraph(doc.subtitle)
    if doc.draft:
        p = d.add_paragraph()
        r = p.add_run("ПРЕДВАРИТЕЛЬНЫЙ ДОКУМЕНТ — окончательная карта уставок не может быть выдана. Причины:\n" + "\n".join("• " + x for x in doc.draft_reasons[:12]))
        r.bold = True; r.font.color.rgb = RGBColor(0xB9, 0x1C, 0x1C)
    if doc.meta:
        t = d.add_table(rows=0, cols=2); t.style = "Table Grid"
        for k, v in doc.meta:
            c = t.add_row().cells
            c[0].text, c[1].text = str(k), str(v)
            c[0].paragraphs[0].runs[0].bold = True
            _shade(c[0], "EEF3F8")
    for b in doc.blocks:
        if b.kind in ("h1", "h2", "h3"):
            d.add_heading(b.text, int(b.kind[1]))
        elif b.kind == "p":
            d.add_paragraph(b.text)
        elif b.kind == "note":
            p = d.add_paragraph(); r = p.add_run("Внимание: " + b.text); r.italic = True
        elif b.kind == "pagebreak":
            d.add_page_break()
        elif b.kind == "image":
            d.add_picture(io.BytesIO(b.data), width=Cm(18))
            d.add_paragraph(b.caption).runs[0].italic = True
        elif b.kind == "table" and b.header:
            if b.caption:
                d.add_paragraph(b.caption)
            t = d.add_table(rows=1, cols=len(b.header)); t.style = "Table Grid"
            for i, h in enumerate(b.header):
                t.rows[0].cells[i].text = str(h)
                t.rows[0].cells[i].paragraphs[0].runs[0].bold = True
                t.rows[0].cells[i].paragraphs[0].runs[0].font.size = Pt(8)
                _shade(t.rows[0].cells[i], "D9E2EC")
            for ri, row in enumerate(b.rows):
                cells = t.add_row().cells
                for i, v in enumerate(row):
                    cells[i].text = str(v)
                    for pr in cells[i].paragraphs:
                        for rn in pr.runs:
                            rn.font.size = Pt(8)
                    s = b.row_status[ri] if ri < len(b.row_status) else None
                    if s in STATUS_COLORS:
                        _shade(cells[i], STATUS_COLORS[s])
            d.add_paragraph()
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ───────────────────────────── XLSX ─────────────────────────────
def to_xlsx(doc: Doc) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Титул"
    ws["A1"] = doc.title; ws["A1"].font = Font(bold=True, size=16, color="0F2A43")
    ws["A2"] = doc.subtitle
    r = 4
    if doc.draft:
        ws.cell(r, 1, "ПРЕДВАРИТЕЛЬНЫЙ ДОКУМЕНТ — не для выдачи. Причины:").font = Font(bold=True, color="B91C1C")
        r += 1
        for x in doc.draft_reasons[:15]:
            ws.cell(r, 1, "• " + x); r += 1
        r += 1
    for k, v in doc.meta:
        ws.cell(r, 1, k).font = Font(bold=True)
        ws.cell(r, 2, str(v)); r += 1
    ws.column_dimensions["A"].width = 30; ws.column_dimensions["B"].width = 100
    thin = Side(style="thin", color="C8CED6")
    used = {"Титул"}
    heading = "Таблица"
    n = 0
    for b in doc.blocks:
        if b.kind in ("h1", "h2", "h3"):
            heading = b.text
        elif b.kind == "table" and b.header:
            n += 1
            name = re.sub(r"[\\/*?:\[\]]", " ", f"{n}. {heading}")[:28]
            base, k = name, 1
            while name in used:
                k += 1; name = f"{base[:25]}~{k}"
            used.add(name)
            s = wb.create_sheet(name)
            s.cell(1, 1, heading).font = Font(bold=True, size=12)
            for i, h in enumerate(b.header):
                c = s.cell(3, i + 1, h)
                c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0F2A43"); c.alignment = Alignment(wrap_text=True, vertical="top")
            for ri, row in enumerate(b.rows):
                st = b.row_status[ri] if ri < len(b.row_status) else None
                for i, v in enumerate(row):
                    c = s.cell(4 + ri, i + 1, v)
                    c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
                    if st in STATUS_COLORS:
                        c.fill = PatternFill("solid", fgColor=STATUS_COLORS[st])
            for i, h in enumerate(b.header):
                mx = max([len(str(h))] + [len(str(rw[i])) for rw in b.rows[:200] if i < len(rw)])
                s.column_dimensions[get_column_letter(i + 1)].width = min(max(10, mx * 0.95), 60)
            s.freeze_panes = "A4"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ───────────────────────────── CSV / JSON ─────────────────────────────
def to_csv(doc: Doc) -> bytes:
    tables = [b for b in doc.blocks if b.kind == "table" and b.header]
    if not tables:
        return b""
    b = max(tables, key=lambda t: len(t.rows))
    out = io.StringIO()
    w = csv.writer(out, delimiter=";", lineterminator="\n")
    w.writerow(b.header)
    for r in b.rows:
        w.writerow([str(x).replace(".", ",") if isinstance(x, float) else x for x in r])
    return ("\ufeff" + out.getvalue()).encode("utf-8")


def to_json(doc: Doc) -> bytes:
    data = {"title": doc.title, "subtitle": doc.subtitle, "meta": doc.meta, "draft": doc.draft, "draft_reasons": doc.draft_reasons,
            "tables": [{"header": b.header, "rows": b.rows, "row_status": b.row_status} for b in doc.blocks if b.kind == "table"]}
    return json.dumps(data, ensure_ascii=False, indent=1).encode("utf-8")


RENDERERS = {"pdf": (to_pdf, "application/pdf"), "docx": (to_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
             "xlsx": (to_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"), "csv": (to_csv, "text/csv; charset=utf-8"), "json": (to_json, "application/json")}
